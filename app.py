import os
import io
import sqlite3
from datetime import date, datetime
from flask import Flask, render_template, request, redirect, url_for, flash, jsonify, send_file

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "dev-secret-change-me")

DB_PATH = os.environ.get("DB_PATH", os.path.join(os.path.dirname(__file__), "worktracker.db"))

# If DATABASE_URL is set (e.g. Render's free PostgreSQL), use Postgres.
# Otherwise fall back to a local SQLite file — this keeps local development
# and Windows testing exactly as before, with zero extra setup.
DATABASE_URL = os.environ.get("DATABASE_URL", "").strip()
USE_POSTGRES = bool(DATABASE_URL)

if USE_POSTGRES:
    import psycopg2
    import psycopg2.extras

STATUS_OPTIONS = ["Yet to Start", "WIP", "Done"]

# Exact column contract for import/export. Order and names are locked —
# an uploaded file must match this exactly, no more, no fewer.
EXPECTED_COLUMNS = ["Name", "Summary", "Details", "Date", "Status"]


class PGConnection:
    """Thin wrapper so Postgres can be used with the same conn.execute(...)
    style calls as sqlite3.Connection, without touching every route."""

    def __init__(self, raw_conn):
        self._conn = raw_conn

    def execute(self, query, params=()):
        cur = self._conn.cursor()
        cur.execute(query.replace("?", "%s"), params)
        return cur

    def commit(self):
        self._conn.commit()

    def close(self):
        self._conn.close()


def get_db():
    if USE_POSTGRES:
        url = DATABASE_URL
        if url.startswith("postgres://"):
            url = url.replace("postgres://", "postgresql://", 1)
        raw_conn = psycopg2.connect(url, cursor_factory=psycopg2.extras.RealDictCursor)
        return PGConnection(raw_conn)
    else:
        conn = sqlite3.connect(DB_PATH)
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA foreign_keys = ON")
        return conn


def init_db():
    conn = get_db()
    if USE_POSTGRES:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS tasks (
                id SERIAL PRIMARY KEY,
                name TEXT NOT NULL,
                summary TEXT,
                details TEXT,
                task_date TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'Yet to Start',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            )
            """
        )
    else:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS tasks (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                summary TEXT,
                details TEXT,
                task_date TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'Yet to Start',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            )
            """
        )
    conn.commit()
    conn.close()


# Create the table on startup — works whether this module is imported by
# gunicorn (Render) or run directly with `python app.py` (local/Windows).
init_db()


def row_to_dict(row):
    return {
        "id": row["id"],
        "name": row["name"],
        "summary": row["summary"],
        "details": row["details"],
        "task_date": row["task_date"],
        "status": row["status"],
        "created_at": row["created_at"],
        "updated_at": row["updated_at"],
    }


def normalize_status(value):
    """Match a status value to a canonical STATUS_OPTIONS entry, case/space-insensitive."""
    if value is None:
        return None
    v = str(value).strip().lower().replace("_", " ")
    for s in STATUS_OPTIONS:
        if s.lower() == v:
            return s
    aliases = {
        "not started": "Yet to Start",
        "todo": "Yet to Start",
        "to do": "Yet to Start",
        "pending": "Yet to Start",
        "in progress": "WIP",
        "inprogress": "WIP",
        "completed": "Done",
        "complete": "Done",
        "finished": "Done",
    }
    return aliases.get(v)


def normalize_date(value):
    """Return an ISO date string from a cell that may be a datetime, date, or text."""
    if value is None or str(value).strip() == "":
        return date.today().isoformat()
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return text  # store as-is if unrecognized; user can fix later


# ---------------------------------------------------------------------------
# Core pages
# ---------------------------------------------------------------------------

@app.route("/")
def index():
    status_filter = request.args.get("status", "All")
    search = request.args.get("q", "").strip()
    sort = request.args.get("sort", "date_desc")
    date_filter = request.args.get("date", "").strip()
    month_filter = request.args.get("month", "").strip()

    query = "SELECT * FROM tasks WHERE 1=1"
    params = []

    if status_filter in STATUS_OPTIONS:
        query += " AND status = ?"
        params.append(status_filter)

    if search:
        query += " AND (name LIKE ? OR summary LIKE ? OR details LIKE ?)"
        like = f"%{search}%"
        params.extend([like, like, like])

    if date_filter:
        query += " AND task_date = ?"
        params.append(date_filter)
    elif month_filter:
        query += " AND task_date LIKE ?"
        params.append(f"{month_filter}-%")

    sort_map = {
        "date_desc": "task_date DESC, id DESC",
        "date_asc": "task_date ASC, id ASC",
        "name_asc": "name COLLATE NOCASE ASC",
        "status": "status ASC, task_date DESC",
    }
    query += f" ORDER BY {sort_map.get(sort, sort_map['date_desc'])}"

    conn = get_db()
    tasks = [row_to_dict(r) for r in conn.execute(query, params).fetchall()]

    counts = {
        "All": conn.execute("SELECT COUNT(*) c FROM tasks").fetchone()["c"],
    }
    for s in STATUS_OPTIONS:
        counts[s] = conn.execute(
            "SELECT COUNT(*) c FROM tasks WHERE status = ?", (s,)
        ).fetchone()["c"]
    conn.close()

    return render_template(
        "index.html",
        tasks=tasks,
        statuses=STATUS_OPTIONS,
        status_filter=status_filter,
        search=search,
        sort=sort,
        date_filter=date_filter,
        month_filter=month_filter,
        counts=counts,
        today=date.today().isoformat(),
    )


@app.route("/task/add", methods=["POST"])
def add_task():
    name = request.form.get("name", "").strip()
    summary = request.form.get("summary", "").strip()
    details = request.form.get("details", "").strip()
    task_date = request.form.get("task_date") or date.today().isoformat()
    status = request.form.get("status", "Yet to Start")

    if not name:
        flash("Task name is required.", "error")
        return redirect(url_for("index"))

    if status not in STATUS_OPTIONS:
        status = "Yet to Start"

    now = datetime.now().isoformat(timespec="seconds")
    conn = get_db()
    conn.execute(
        """INSERT INTO tasks (name, summary, details, task_date, status, created_at, updated_at)
           VALUES (?, ?, ?, ?, ?, ?, ?)""",
        (name, summary, details, task_date, status, now, now),
    )
    conn.commit()
    conn.close()
    flash("Task added.", "success")
    return redirect(url_for("index"))


@app.route("/task/<int:task_id>/edit", methods=["POST"])
def edit_task(task_id):
    name = request.form.get("name", "").strip()
    summary = request.form.get("summary", "").strip()
    details = request.form.get("details", "").strip()
    task_date = request.form.get("task_date")
    status = request.form.get("status", "Yet to Start")

    if not name:
        flash("Task name is required.", "error")
        return redirect(url_for("index"))

    if status not in STATUS_OPTIONS:
        status = "Yet to Start"

    now = datetime.now().isoformat(timespec="seconds")
    conn = get_db()
    conn.execute(
        """UPDATE tasks SET name=?, summary=?, details=?, task_date=?, status=?, updated_at=?
           WHERE id=?""",
        (name, summary, details, task_date, status, now, task_id),
    )
    conn.commit()
    conn.close()
    flash("Task updated.", "success")
    return redirect(url_for("index"))


@app.route("/task/<int:task_id>/status", methods=["POST"])
def quick_status(task_id):
    status = request.form.get("status")
    if status not in STATUS_OPTIONS:
        return jsonify({"ok": False, "error": "invalid status"}), 400
    now = datetime.now().isoformat(timespec="seconds")
    conn = get_db()
    conn.execute("UPDATE tasks SET status=?, updated_at=? WHERE id=?", (status, now, task_id))
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/task/<int:task_id>/delete", methods=["POST"])
def delete_task(task_id):
    conn = get_db()
    conn.execute("DELETE FROM tasks WHERE id=?", (task_id,))
    conn.commit()
    conn.close()
    flash("Task deleted.", "success")
    return redirect(url_for("index"))


@app.route("/health")
def health():
    return jsonify({"status": "ok"})


# ---------------------------------------------------------------------------
# Excel: styling helpers
# ---------------------------------------------------------------------------

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
EXAMPLE_FONT = Font(name="Calibri", size=11, italic=True, color="808080")
BODY_FONT = Font(name="Calibri", size=11, color="1A1F2B")
THIN_BORDER = Border(
    left=Side(style="thin", color="D9DEE6"),
    right=Side(style="thin", color="D9DEE6"),
    top=Side(style="thin", color="D9DEE6"),
    bottom=Side(style="thin", color="D9DEE6"),
)
COLUMN_WIDTHS = [30, 36, 50, 14, 16]


def style_header_row(ws):
    for idx, col_name in enumerate(EXPECTED_COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
        cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(idx)].width = COLUMN_WIDTHS[idx - 1]
    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"


# ---------------------------------------------------------------------------
# Excel: sample template
# ---------------------------------------------------------------------------

@app.route("/template/download")
def download_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Work Tracker Template"
    style_header_row(ws)

    example = ["Fix ORA-00060 deadlock", "Investigate prod deadlock",
               "Checked AWR report, found blocking session, coordinated kill with DBA team",
               date.today().isoformat(), "WIP"]
    for idx, value in enumerate(example, start=1):
        cell = ws.cell(row=2, column=idx, value=value)
        cell.font = EXAMPLE_FONT
        cell.border = THIN_BORDER

    ws.cell(row=1, column=1).comment = Comment(
        "Do not rename, reorder, remove, or add columns. "
        "The import only accepts exactly these 5 headers: "
        + ", ".join(EXPECTED_COLUMNS) + ".",
        "Work Tracker",
    )

    note_row = 4
    ws.cell(row=note_row, column=1,
             value="Row 2 is a sample — delete it before importing your own data.").font = Font(
        italic=True, size=10, color="808080"
    )
    ws.cell(row=note_row + 1, column=1,
             value="Status must be one of: " + ", ".join(STATUS_OPTIONS)).font = Font(
        italic=True, size=10, color="808080"
    )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name="work_tracker_template.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# ---------------------------------------------------------------------------
# Excel: export current data
# ---------------------------------------------------------------------------

@app.route("/export/excel")
def export_excel():
    conn = get_db()
    rows = conn.execute("SELECT * FROM tasks ORDER BY task_date DESC, id DESC").fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Work Tracker Export"
    style_header_row(ws)

    for r_idx, row in enumerate(rows, start=2):
        values = [row["name"], row["summary"], row["details"], row["task_date"], row["status"]]
        for c_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.font = BODY_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(c_idx == 3))

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    stamp = date.today().isoformat()
    return send_file(
        buf,
        as_attachment=True,
        download_name=f"work_tracker_export_{stamp}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# ---------------------------------------------------------------------------
# Excel: import with strict column validation
# ---------------------------------------------------------------------------

@app.route("/import/excel", methods=["POST"])
def import_excel():
    file = request.files.get("file")

    if not file or file.filename == "":
        flash("Please choose an .xlsx file to import.", "error")
        return redirect(url_for("index"))

    if not file.filename.lower().endswith((".xlsx", ".xlsm")):
        flash("Only .xlsx files are supported. Please use the provided template.", "error")
        return redirect(url_for("index"))

    try:
        wb = load_workbook(file, data_only=True)
    except Exception:
        flash("That file couldn't be read as an Excel workbook. Please use the provided template.", "error")
        return redirect(url_for("index"))

    ws = wb.active
    header_cells = next(ws.iter_rows(min_row=1, max_row=1), [])
    actual_headers = [str(c.value).strip() if c.value is not None else "" for c in header_cells]
    while actual_headers and actual_headers[-1] == "":
        actual_headers.pop()

    if actual_headers != EXPECTED_COLUMNS:
        expected_set = set(EXPECTED_COLUMNS)
        actual_set = set(actual_headers)
        extra = [c for c in actual_headers if c not in expected_set]
        missing = [c for c in EXPECTED_COLUMNS if c not in actual_set]

        parts = []
        if extra:
            parts.append("unexpected column(s): " + ", ".join(f'"{c}"' for c in extra))
        if missing:
            parts.append("missing column(s): " + ", ".join(f'"{c}"' for c in missing))
        if not parts and actual_headers != EXPECTED_COLUMNS:
            parts.append("columns are not in the required order: " + ", ".join(EXPECTED_COLUMNS))

        flash(
            "Import rejected — " + "; ".join(parts) + ". "
            "The file must use exactly these columns, in this order: "
            + ", ".join(EXPECTED_COLUMNS) + ". Download the template and try again.",
            "error",
        )
        return redirect(url_for("index"))

    now = datetime.now().isoformat(timespec="seconds")
    conn = get_db()
    imported = 0
    skipped_rows = []

    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue

        name = (str(row[0]).strip() if row[0] is not None else "")
        summary = (str(row[1]).strip() if len(row) > 1 and row[1] is not None else "")
        details = (str(row[2]).strip() if len(row) > 2 and row[2] is not None else "")
        raw_date = row[3] if len(row) > 3 else None
        raw_status = row[4] if len(row) > 4 else None

        if not name:
            skipped_rows.append(row_num)
            continue

        task_date = normalize_date(raw_date)
        status = normalize_status(raw_status) or "Yet to Start"

        conn.execute(
            """INSERT INTO tasks (name, summary, details, task_date, status, created_at, updated_at)
               VALUES (?, ?, ?, ?, ?, ?, ?)""",
            (name, summary, details, task_date, status, now, now),
        )
        imported += 1

    conn.commit()
    conn.close()

    if imported == 0 and not skipped_rows:
        flash("No data rows found below the header — nothing was imported.", "error")
    else:
        msg = f"Imported {imported} task(s) from {file.filename}."
        if skipped_rows:
            msg += f" Skipped {len(skipped_rows)} row(s) with no task name (row(s): {', '.join(map(str, skipped_rows))})."
        flash(msg, "success")

    return redirect(url_for("index"))


if __name__ == "__main__":
    init_db()
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=True)